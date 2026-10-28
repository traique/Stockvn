"""Tài sản: the owner's real-money journal (vnbot.taisan), the Excel import/export, the /ts
commands, the /taisan web tab and the 15:30 report (offline, prices from a fake market)."""
import io
import json
import os
import tempfile
import time
from datetime import datetime
from pathlib import Path

import _common  # noqa: F401 - clean env, no .env, offline, temp folder, repo on sys.path
_common.need("httpx")
_common.need("openpyxl")

ROOT = Path(tempfile.mkdtemp(prefix="taisan-"))
os.environ.update({"DATA_DIR": str(ROOT / "app"), "WEB_PASSWORD": "s3cret-pass", "TRADE_MODE": "paper",
                   "PUBLIC_URL": "https://bot.example", "TRADINGAGENTS_PARAMS_PATH": str(ROOT / "qp.json"),
                   "TELEGRAM_TOKEN": "123456789:AAHdqTcvCH1vGWJxfSeofSAs0K5PALDsawQ", "TELEGRAM_ALLOWED_IDS": "111"})

from vnbot import settings as settings_module  # noqa: E402

SETTINGS = settings_module.load()

import openpyxl  # noqa: E402
from tradingagents.market import vn_rules  # noqa: E402
from vnbot import taisan as ts  # noqa: E402
from vnbot import taisan_excel as tx  # noqa: E402
from vnbot.taisan import TaiSan  # noqa: E402

t = _common.Checks()
NOW = {"t": datetime(2026, 10, 24, 10, 0, tzinfo=vn_rules.VN_TZ)}       # Saturday: no live prices
vn_rules.now_vn = lambda: NOW["t"]


def at(day, hhmm):
    h, m = map(int, hhmm.split(":"))
    y, mo, d = map(int, day.split("-"))
    return datetime(y, mo, d, h, m, tzinfo=vn_rules.VN_TZ)


SESSIONS = []
_d = datetime(2026, 3, 2).date()
while _d.isoformat() <= "2026-10-30":
    if vn_rules.is_session_day(_d.isoformat()):
        SESSIONS.append(_d.isoformat())
    _d = _d.fromordinal(_d.toordinal() + 1)

BASE = {"VNINDEX": 1250.0, "FPT": 30_000.0, "HPG": 25_000.0, "CTS": 20_000.0, "VNM": 60_000.0}


class FakeMarket:
    """Closes that rise 0.1% a session from a base, overrides per (symbol, day), events."""

    def __init__(self, last="2026-10-23"):
        self.last, self.over, self.ev, self.calls = last, {}, {}, []

    def px(self, symbol, day):
        if (symbol, day) in self.over:
            return self.over[(symbol, day)]
        i = SESSIONS.index(day)
        return round(BASE.get(symbol, 10_000.0) * (1 + 0.001 * i), 2)

    def daily_closes(self, symbol, end):
        self.calls.append(("daily", symbol, end))
        return {d: self.px(symbol, d) for d in SESSIONS if d <= min(end, self.last)}

    def unadjusted_closes(self, symbol, start, end):
        self.calls.append(("unadj", symbol, start, end))
        return {d: v for d, v in self.daily_closes(symbol, end).items() if d >= start}

    def live(self, symbols, day):
        self.calls.append(("live", tuple(symbols), day))
        return {s: self.px(s, day) for s in symbols if day in SESSIONS}

    def events(self, symbol, start, end):
        return [e for e in self.ev.get(symbol, []) if start <= e["exrightDate"][:10] <= end]

    def exchange(self, symbol):
        return "HOSE"


def account(name, market=None):
    return TaiSan(ROOT / name, market=market or FakeMarket(), site=None)


def raises(fn, needle=""):
    try:
        fn()
    except ValueError as exc:
        return needle in str(exc)
    return False


# ───────────── 1. the journal: what each entry does ─────────────
a = account("core")
a.add("DEPOSIT", "2026-04-13", amount=100_000_000, note="vốn")
e, b = a.add("BUY", "2026-04-14", symbol="fpt", qty=1000, price=30_000)
t.ok("buy: symbol upper-cased, cash paid, cost kept", e.symbol == "FPT" and b.cash == 70_000_000
     and b.holdings["FPT"].avg == 30_000, (e.symbol, b.cash))
t.ok("T+1: a sale refused (shares not arrived)",
     raises(lambda: a.add("SELL", "2026-04-15", symbol="FPT", qty=500, price=31_000), "mới có 0 cp"))
h = a.book().holdings["FPT"]
t.ok("T+2: sellable from the afternoon, not the morning",
     h.sellable("2026-04-16", afternoon=False) == 0 and h.sellable("2026-04-16") == 1000)
t.ok("pending lots name the T+2 day", h.pending_lots("2026-04-15") == [("2026-04-16", 1000)], h.pending_lots("2026-04-15"))
e, b = a.add("SELL", "2026-04-16", symbol="FPT", qty=400, price=33_000)
row = next(r for r in b.rows if r.entry.id == e.id)
t.ok("sell: P&L at the average cost, as entered (no fee/tax)", row.pnl == 1_200_000 and b.realized == 1_200_000
     and b.cash == 83_200_000 and b.wins == 1, (row.pnl, b.cash))
e, b = a.add("BONUS", "2026-04-17", symbol="FPT", qty=60)
h = b.holdings["FPT"]
t.ok("bonus shares: qty up, cost unchanged, avg down", h.qty == 660 and h.cost == 18_000_000
     and abs(h.avg - 27_272.727) < 0.01, (h.qty, h.cost))
t.ok("bonus shares sellable at once", h.sellable("2026-04-17", afternoon=False) == 660)
a.add("DIVIDEND", "2026-05-04", symbol="FPT", amount=1_000_000)
a.add("WITHDRAW", "2026-05-04", amount=4_200_000)
e, b = a.add("ADJUST", "2026-05-04", amount=-100_000, note="phí lưu ký")
t.ok("cash kinds: dividend in, withdraw out (capital down), adjust in P&L only",
     b.cash == 79_900_000 and b.dividends == 1_000_000 and b.capital == 95_800_000 and b.adjustments == -100_000,
     (b.cash, b.capital))
# Same day: the sale is booked before the buy, at the average cost before that buy.
a.add("BUY", "2026-06-01", symbol="FPT", qty=100, price=40_000)
e, b = a.add("SELL", "2026-06-01", symbol="FPT", qty=300, price=31_000)
row = next(r for r in b.rows if r.entry.id == e.id)
t.ok("same day: sell before buy, P&L at the old average", abs(row.avg_cost - 27_272.727) < 0.01
     and abs(row.pnl - 1_118_181.82) < 0.01, (row.avg_cost, row.pnl))
t.ok("cash after both", abs(b.cash - 85_200_000) < 0.5 and b.holdings["FPT"].qty == 460, b.cash)
t.ok("refused: more than held", raises(lambda: a.add("SELL", "2026-06-02", symbol="FPT", qty=1000, price=31_000), "chỉ có 460"))
t.ok("refused: a day in the future", raises(lambda: a.add("DEPOSIT", "2026-10-26", amount=1), "tương lai"))
t.ok("refused: a trade on a weekend", raises(lambda: a.add("BUY", "2026-10-24", symbol="FPT", qty=1, price=30_000),
                                             "không có phiên"))
t.ok("refused: bad symbol", raises(lambda: a.add("BUY", "2026-06-02", symbol="FP T!", qty=1, price=30_000), "Mã"))
t.ok("refused: fractional shares", raises(lambda: a.add("BUY", "2026-06-02", symbol="FPT", qty=1.5, price=30_000), "nguyên"))
t.ok("refused: a price in thousands", raises(lambda: a.add("BUY", "2026-06-02", symbol="FPT", qty=1, price=30.5), "Giá"))
t.ok("refused: zero deposit, unknown kind", raises(lambda: a.add("DEPOSIT", "2026-06-02", amount=0))
     and raises(lambda: a.add("XYZ", "2026-06-02")))
n_before = len(a.entries())
e, b = a.add("WITHDRAW", "2026-06-02", amount=90_000_000)
t.ok("negative cash: recorded with a warning", b.cash < 0 and any("Tiền mặt âm" in w for w in b.warnings), b.warnings)
t.ok("undo the last line", a.undo(source="web").startswith("↩️ Đã huỷ #") and a.book().cash > 0
     and not a.book().warnings and len(a.entries()) == n_before + 2)
first_buy = next(x for x in a.entries() if x.kind == "BUY")
t.ok("undo refused when later lines need it", raises(lambda: a.void(first_buy.id), "Không huỷ được"))
t.ok("undo of an undone line refused", raises(lambda: a.void(e.id), "đã huỷ"))
with open(a.journal_path, "ab") as f:
    f.write(b'{"id": 999, "kind": "BU')                  # a line cut by a crash
a.add("DEPOSIT", "2026-06-02", amount=1_000_000)
t.ok("a cut line is skipped, the next one still reads", len(a.entries()) == n_before + 3
     and a.entries()[-1].kind == "DEPOSIT", [x.kind for x in a.entries()[-3:]])
rows = ts.derive(a.entries()).rows
t.ok("journal text: newest first, P&L of sales, undone mark", "#3 16/04/2026 Bán FPT 400 cp @ 33.000 · lãi/lỗ +1.200.000"
     in a.journal_text(30) and "ĐÃ HUỶ" in a.journal_text(30), a.journal_text(30)[:400])
t.ok("derive upto a day", ts.derive(a.entries(), upto="2026-04-14").cash == 70_000_000 and len(rows) >= 9)

# ───────────── 2. the asset history: fund units, VNINDEX ─────────────
flat = FakeMarket()
flat.px = lambda symbol, day: (1000.0 if symbol == "VNINDEX" else 10_000.0) * (1.1 if day >= "2026-05-04" and symbol == "AAA"
                                                                                      else 1.0)
h1 = account("hist", flat)
h1.add("DEPOSIT", "2026-04-13", amount=10_000_000)
h1.add("BUY", "2026-04-14", symbol="AAA", qty=1000, price=10_000)
h1.add("DEPOSIT", "2026-04-20", amount=5_000_000)          # money in: not a gain
h1.add("WITHDRAW", "2026-04-22", amount=2_000_000)         # money out: not a loss
h1.rebuild(through="2026-05-05")
pts = h1.history()
by = {p["day"]: p for p in pts}
t.ok("history from the first entry, one point per session", pts[0]["day"] == "2026-04-13" and "2026-04-18" not in by
     and "2026-05-05" in by, [p["day"] for p in pts[:3]])
t.ok("deposit and withdrawal leave the unit value at 100", all(abs(by[d]["nav"] - 100) < 1e-6 for d in
                                                                ("2026-04-14", "2026-04-20", "2026-04-22", "2026-04-29")),
     [(d, by[d]["nav"]) for d in ("2026-04-20", "2026-04-22")])
t.ok("a 10% rise of a fully invested book: +10%... of the stock part", abs(by["2026-05-04"]["nav"] - 100 * (13_000_000 + 1_000_000)
                                                                          / 13_000_000) < 1e-3, by["2026-05-04"])
t.ok("total = cash + stock value", by["2026-05-04"]["total"] == 3_000_000 + 11_000_000 and by["2026-05-04"]["capital"]
     == 13_000_000, by["2026-05-04"])
t.ok("VNINDEX indexed to 100 on the first point", pts[0]["vni_index"] == 100 and pts[-1]["vni_index"] == 100)
t.ok("rebuilt flag up to the rebuild day", by["2026-05-05"]["rebuilt"] and h1.state()["rebuilt_until"] == "2026-05-05")
h1.void(next(x.id for x in h1.entries() if x.kind == "WITHDRAW"))
t.ok("an undo shows in the history at once", h1.history()[-1]["capital"] == 15_000_000)

# ───────────── 3. rights waiting to arrive ─────────────
mk = FakeMarket()
r = account("rights", mk)
r.add("DEPOSIT", "2026-06-01", amount=200_000_000)
r.add("BUY", "2026-06-01", symbol="HPG", qty=3500, price=25_000)
mk.ev["HPG"] = [
    {"eventCode": "DIV", "exrightDate": "2026-07-01", "valuePerShare": 1000, "eventTitleVi": "Trả cổ tức bằng tiền"},
    {"eventCode": "ISS", "exrightDate": "2026-07-15", "exerciseRatio": 0.1, "eventTitleVi": "Cổ phiếu thưởng 10%"},
    {"eventCode": "ISS", "exrightDate": "2026-08-03", "exerciseRatio": 0.5, "eventTitleVi": "Phát hành quyền mua"},
    {"eventCode": "DIV", "exrightDate": "2026-10-28", "valuePerShare": 500, "eventTitleVi": "Cổ tức đợt 2"},
]
items = r.refresh_events("2026-10-23")
cash_item = next((p for p in items if p["kind"] == "CASH"), None)
stock_item = next((p for p in items if p["kind"] == "STOCK"), None)
t.ok("cash dividend waiting: after 5% tax", cash_item and cash_item["value"] == 3_325_000 and cash_item["eligible"] == 3500,
     cash_item)
t.ok("bonus shares waiting: 10% of the shares held", stock_item and stock_item["shares"] == 350
     and stock_item["value"] == round(350 * mk.px("HPG", "2026-10-23")), stock_item)
t.ok("rights to buy are not counted", len(items) == 2, items)
t.ok("a coming ex-date is listed", [u["ex"] for u in r.upcoming_rights()] == ["2026-10-28"], r.upcoming_rights())
s = r.summary()
t.ok("rights waiting count in the total", abs(s["pending"] - (3_325_000 + stock_item["value"])) < 1
     and abs(s["total"] - (s["cash"] + s["market"] + s["pending"])) < 1, s)
r.handle("/tscotuc", "HPG 950")                             # 950 a share after tax, 3,500 from the waiting right
t.ok("/tscotuc takes the eligible shares of the waiting right", r.book().dividends == 3_325_000, r.book().dividends)
r.handle("/tsthuong", "HPG 350")
items = r.refresh_events("2026-10-23")
t.ok("recorded rights leave the waiting list", items == [], items)
t.ok("pending value per day kept for the chart", ts._read_json(r.pending_path, {})["by_day"].get("2026-10-23") == 0)

# ───────────── 4. alerts and their shadow results ─────────────
t.ok("alert kept once per day/symbol/level", r.record_alert("stop", "HPG", 24_000, "2026-10-01", "10:15")
     and not r.record_alert("stop", "HPG", 23_900, "2026-10-01") and r.record_alert("sell", "HPG", 24_000, "2026-10-01"))
t.ok("unknown level ignored", not r.record_alert("xyz", "HPG", 1, "2026-10-01"))
done = r.settle_shadow("2026-10-23")
rows = r.state()["shadow"]
i = SESSIONS.index("2026-10-01")
want5 = round((mk.px("HPG", SESSIONS[i + 5]) / mk.px("HPG", "2026-10-01") - 1) * 100, 2)
t.ok("shadow results 5 and 10 sessions later", done == 2 and rows[0]["r5"] == want5 and rows[0]["r10"] is not None, rows)
st = {x["level"]: x for x in r.shadow_stats()}
t.ok("shadow stats per level", st["stop"]["settled"] == 1 and st["stop"]["right"] == 0 and st["early"]["alerts"] == 0, st)
t.ok("settled rows not asked again", r.settle_shadow("2026-10-23") == 0)

# ───────────── 5. what to do, allocation, the report ─────────────
r.add("BUY", "2026-10-21", symbol="FPT", qty=100, price=30_000)
r.set_thresholds(20, 30)
r.set_levels("HPG", 99_000, None)          # the close is under the cut-loss
r.set_levels("FPT", 20_000, 25_000)        # the close is over the target
r.rebuild(through="2026-10-22")
r.record_day("2026-10-23")
s = r.summary()
todo = [line for _, line in r.actions("2026-10-23", s, {p["symbol"]: p["price"] for p in s["positions"]}, {})]
t.ok("⛔ under the cut-loss first, sellable shares said", todo and todo[0].startswith("⛔ HPG")
     and "bán được 3.850 cp" in todo[0], todo)
t.ok("🎯 target reached; FPT bought Wed sells from Fri afternoon", any(x.startswith("🎯 FPT") for x in todo), todo)
t.ok("⚠️ over the danger weight", any(x.startswith("⚠️ HPG") for x in todo), todo)
alloc = r.allocation(s)
t.ok("allocation levels and sectors", alloc["symbols"][0]["symbol"] == "HPG" and alloc["symbols"][0]["level"] == "danger"
     and alloc["sectors"], alloc)
t.ok("thresholds checked", raises(lambda: r.set_thresholds(40, 20)) and r.thresholds() == (20.0, 30.0))
report = r.report_text("2026-10-23", refresh_site=False)
for piece in ("💼 TÀI SẢN · phiên 23/10/2026", "📋 VIỆC CẦN LÀM PHIÊN 26/10", "⛔ HPG", "📈 TỪNG MÃ", "⚖️ PHÂN BỔ",
              "🗓️ HPG: cổ tức tiền 500 đ/cp", "🧪 KẾT QUẢ CẢNH BÁO", "Bot không đặt lệnh"):
    t.ok(f"report has {piece!r}", piece in report, report)
t.ok("no fee or tax taken off anywhere", "phí" not in report.split("Giá theo đúng")[0].lower(), report)

# ───────────── 6. the commands ─────────────
c = account("cmd")
say = c.handle
t.ok("/ts empty", "chưa có dữ liệu" in say("/ts", ""))
t.ok("/tsnap", say("/tsnap", "50.000.000 2026-04-13 vốn đầu").startswith("✅ #1 Nạp tiền 50.000.000 ngày 13/04/2026"))
reply = say("/tsmua", "fpt 1000 30.000 2026-04-14 lô đầu")
t.ok("/tsmua: T+2 day, cash, no level hint", "MUA FPT 1.000 cp @ 30.000" in reply and "Bán được từ chiều 16/04" in reply
     and "Tiền mặt còn 20.000.000" in reply and "/tscatlo FPT" in reply, reply)
t.ok("/tsban refused before T+2", raises(lambda: say("/tsban", "FPT 100 31000 2026-04-15"), "T+2"))
reply = say("/tsban", "FPT 400 33000 2026-04-16")
t.ok("/tsban: P&L and what is left", "lãi/lỗ +1.200.000 (+10,00%)" in reply and "Còn 600 cp" in reply, reply)
t.ok("/tsthuong", "+60 cp thưởng" in say("/tsthuong", "FPT 60 2026-04-17"))
reply = say("/tscotuc", "FPT 2.000")
t.ok("/tscotuc: per share x shares held", "2.000 đ/cp x 660 cp" in reply and c.book().dividends == 1_320_000, reply)
t.ok("/tsrut", "Rút tiền 1.000.000" in say("/tsrut", "1000000"))
reply = say("/tstien", "34.000.000")
t.ok("/tstien: one adjustment to the broker's balance", "Điều chỉnh tiền mặt +480.000" in reply
     and abs(c.book().cash - 34_000_000) < 1 and c.book().capital == 49_000_000, reply)
t.ok("/tstien again: nothing to do", "không cần điều chỉnh" in say("/tstien", "34000000"))
t.ok("/tscatlo", "cắt lỗ 25.000, mục tiêu 40.000" in say("/tscatlo", "FPT 25000 40000")
     and c.levels()["FPT"]["source"] == "owner")
t.ok("/tscatlo refused for a stock not held", raises(lambda: say("/tscatlo", "VNM 1000"), "không có VNM"))
t.ok("/tscatlo xoa", "Đã xoá" in say("/tscatlo", "FPT xoa") and "FPT" not in c.levels())
t.ok("/tsnguong show and set", "cảnh báo 25%" in say("/tsnguong", "") and "20%" in say("/tsnguong", "20 35")
     and raises(lambda: say("/tsnguong", "a b")))
t.ok("/tsnhatky", "🧾 NHẬT KÝ" in say("/tsnhatky", "5"))
t.ok("/tshuy 7 then the last", say("/tshuy", "7").startswith("↩️ Đã huỷ #7") and say("/tshuy", "").startswith("↩️ Đã huỷ #6"))
t.ok("/ts overview", "💼 TÀI SẢN" in say("/ts", "") and "FPT" in say("/ts", ""))
reply = say("/tsxuat", "")
t.ok("/tsxuat on Telegram: an .xlsx file", isinstance(reply, dict) and Path(reply["file"]).suffix == ".xlsx"
     and Path(reply["file"]).read_bytes()[:2] == b"PK", reply)
reply = say("/tsxuat", "", source="zoom", public_url="https://bot.example")
t.ok("/tsxuat on Zoom: the web link", reply.endswith("https://bot.example/taisan/xuat.xlsx"), reply)
t.ok("/tsnhap without a file says how", "kèm chú thích /tsnhap" in say("/tsnhap", "")
     and "Telegram" in say("/tsnhap", "", source="zoom"))
t.ok("/tsnhap xacnhan with nothing waiting", raises(lambda: say("/tsnhap", "xacnhan"), "Chưa có file"))
t.ok("unknown /ts command: the help", raises(lambda: say("/tsxyz", ""), "/tsmua FPT"))
t.ok("/tsbaocao day: the last session before 15:05", c.report_day() == "2026-10-23")
NOW["t"] = at("2026-10-26", "11:00")
t.ok("/tsbaocao day in a session: the one before", c.report_day() == "2026-10-23")
NOW["t"] = at("2026-10-26", "15:10")
t.ok("/tsbaocao day after the close: today", c.report_day() == "2026-10-26")
NOW["t"] = at("2026-10-24", "10:00")

# ───────────── 7. the bot's view (morning brief, scans, intraday) ─────────────
v = account("view")
v.add("DEPOSIT", "2026-10-19", amount=100_000_000)
v.add("BUY", "2026-10-20", symbol="HPG", qty=1000, price=25_000)
v.add("BUY", "2026-10-22", symbol="HPG", qty=500, price=26_000)
v.set_levels("HPG", 23_000, 28_000)
view = v.ledger_view()
pos = view.positions["HPG"]
t.ok("view: qty, average, levels, cash", pos.qty == 1500 and abs(pos.avg_price - 25_333.33) < 0.01 and pos.stop == 23_000
     and pos.target1 == 28_000 and view.cash == 62_000_000 and view.meta.get("taisan"), (pos, view.cash))
t.ok("view: the recent lots for T+2", sorted(lot["date"] for lot in pos.unsettled) == ["2026-10-20", "2026-10-22"], pos.unsettled)
pos.last_price, pos.last_date, pos.stop = 24_000.0, "2026-10-23", 11_500.0       # an ex-date halved the levels
pos.target1 = 14_000.0
v.save_view(view)
st = v.state()
t.ok("save_view keeps marks and moved levels", st["marks"]["HPG"]["last_price"] == 24_000 and st["levels"]["HPG"]["stop"] == 11_500
     and st["levels"]["HPG"]["target"] == 14_000, st)
prices, label = v.current_prices(["HPG"])
t.ok("price outside the session: the evening mark", prices == {"HPG": 24_000.0} and label == "giá đóng cửa gần nhất", (prices, label))
NOW["t"] = at("2026-10-23", "10:30")
prices, label = v.current_prices(["HPG"])
t.ok("price in the session: live, cached a minute", label == "khớp 10:30" and prices["HPG"] == v.market.px("HPG", "2026-10-23")
     and v.current_prices(["HPG"]) and sum(1 for x in v.market.calls if x[0] == "live") == 1, (prices, label))
NOW["t"] = at("2026-10-24", "10:00")

# ───────────── 8. Excel: the old app's file, this module's own export ─────────────
def old_app_file(cash=58_784_700.0):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vị thế hiện tại"
    ws.append(tx.POSITION_HEADERS)
    ws.append(["FPT", 30000, 600, "2026-04-14", 31000, 18_000_000, 18_600_000, 600_000, 3.33])
    ws.append(["HPG", 22727.27, 1100, "2026-04-14", 23000, 25_000_000, 25_300_000, 300_000, 1.2])
    ws.append(["CTS", 20000, 3947.37, "2026-06-02", 21000, 78_947_400, 82_894_770, 3_947_370, 5])  # lots split in fractions
    ws.append(["CTS", 20000, 1052.63, None, 21000, 21_052_600, 22_105_230, 1_052_630, 5])           # no buy date
    hs = wb.create_sheet("Lịch sử giao dịch")
    hs.append(tx.HISTORY_HEADERS)
    hs.append(["2026-07-01", "Bán", "HPG", 1100, 24000, 26_400_000, 22727.27, 1_400_000, ""])
    hs.append(["2026-06-02", "Mua", "CTS", 5000, 20000, 100_000_000, "", "", ""])
    hs.append(["01/06/2026", "Bán", "HPG", 200, 0, 0, "", "", "thưởng"])                       # bonus written as a sale at 0
    hs.append(["2026-05-04", "Bán", "FPT", 400, 33000, 13_200_000, 30000, 1_200_000, "=cmd|' /C calc'!A0"])
    hs.cell(hs.max_row, 9).data_type = "s"          # a note that looks like a formula, kept as text
    hs.append(["2026-04-14", "Mua", "HPG", 2000, 25000, 50_000_000, "", "", ""])
    hs.append([datetime(2026, 4, 14), "Mua", "FPT", 1000, 30000, 30_000_000, "", "", ""])
    sm = wb.create_sheet("Tổng quan")
    for k, val in (("Vốn gốc", 200_000_000), ("Tiền mặt", cash), ("Lãi/Lỗ chốt", 2_600_000), ("Lãi/Lỗ tổng", 7_684_700)):
        sm.append([k, val])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


x = account("excel")
text = tx.preview(x, old_app_file(), filename="portfolio.xlsx")
t.ok("preview: every check matches the file", text.count("✅") == 5 and "⚠️" not in text, text)
t.ok("preview: the old app's faults fixed and said", "Sửa 1 dòng “Bán … giá 0”" in text and "CTS 5.000 cp" in text
     and "Lỗ trước khi theo dõi (khớp tiền mặt file Excel): -815.300" in text, text)
t.ok("preview writes nothing to the journal", x.entries() == [] and tx.pending(x)["filename"] == "portfolio.xlsx")
reply = tx.confirm(x, replace=False)
bk = x.book()
t.ok("confirm: the book as the file", reply.startswith("✅ Đã nhập 8 dòng") and abs(bk.cash - 58_784_700) < 1
     and bk.capital == 200_000_000 and abs(bk.realized - 2_600_000) < 1
     and {s: h.qty for s, h in bk.holdings.items()} == {"FPT": 600, "HPG": 1100, "CTS": 5000}, (reply, bk.cash))
t.ok("imported lines marked, the pending file gone", all(e.source == "import" for e in x.entries())
     and tx.pending(x) is None and x.state()["imported"]["filename"] == "portfolio.xlsx")
tx.preview(x, old_app_file(), filename="again.xlsx")
t.ok("a second import needs 'thay thế'", raises(lambda: tx.confirm(x, replace=False), "thaythe"))
tx.confirm(x, replace=True)
t.ok("replace keeps the old journal as a copy", len(list(x.dir.glob("journal.jsonl.bak-*"))) == 1 and len(x.entries()) == 8)
tx.preview(x, old_app_file(), filename="b.xlsx")
t.ok("cancel", tx.cancel(x) == "Đã bỏ file chờ nhập." and tx.pending(x) is None and tx.cancel(x) == "Không có file nào chờ nhập.")
for name, data, needle in (("not a workbook", b"hello", "Không phải file Excel"), ("too big", b"PK" + b"0" * 2_000_001, "quá lớn"),
                           ("a broken zip", b"PK\x03\x04garbage", "Không đọc được")):
    t.ok(f"refused: {name}", raises(lambda data=data: tx.preview(x, data), needle))
wb = openpyxl.Workbook()
wb.active.title = "Khác"
buf = io.BytesIO()
wb.save(buf)
t.ok("refused: no history sheet", raises(lambda: tx.preview(x, buf.getvalue()), "thiếu sheet"))
wb = openpyxl.Workbook()
wb.active.title = "Lịch sử giao dịch"
wb.active.append(tx.HISTORY_HEADERS)
wb.active.append(["2026-05-04", "Bán", "FPT", 100, 33000, 3_300_000, "", "", ""])
wb.create_sheet("Tổng quan").append(["Vốn gốc", 10_000_000])
buf = io.BytesIO()
wb.save(buf)
reply = tx.preview(x, buf.getvalue())
t.ok("an impossible history: a clear error, nothing kept", reply.startswith("❌") and "chỉ có 0 cp" in reply
     and tx.pending(x) is None, reply)
# This module's export reads back to the same book (the cash sheet carries every flow).
x.add("DIVIDEND", "2026-07-02", symbol="FPT", amount=1_200_000, note="=1+1")
x.add("WITHDRAW", "2026-07-03", amount=5_000_000)
path = tx.export_xlsx(x, ROOT / "out" / "x.xlsx")
wbx = openpyxl.load_workbook(path)
t.ok("export: the four sheets of the old app", wbx.sheetnames == ["Vị thế hiện tại", "Lịch sử giao dịch", "Tổng quan", "Tiền mặt"])
t.ok("export: a note typed as a formula stays text", any(c.value == "=1+1" and c.data_type == "s"
                                                          for row in wbx["Tiền mặt"].iter_rows() for c in row))
y = account("excel2")
tx.preview(y, path.read_bytes(), filename="x.xlsx")
tx.confirm(y, replace=False)
bx, by_ = x.book(), y.book()
t.ok("round trip: same cash, capital, P&L, dividends, holdings",
     abs(bx.cash - by_.cash) < 1 and bx.capital == by_.capital and abs(bx.realized - by_.realized) < 1
     and bx.dividends == by_.dividends and {s: (h.qty, round(h.cost)) for s, h in bx.holdings.items()}
     == {s: (h.qty, round(h.cost)) for s, h in by_.holdings.items()}, (bx.cash, by_.cash))
csv_text = tx.export_csv(x)
t.ok("CSV: a BOM for Excel, formula-like notes neutralised", csv_text.startswith("﻿Ngày,Loại")
     and "'=1+1" in csv_text and "'=cmd" in csv_text, csv_text[:300])

# ───────────── 9. the web tab ─────────────
from fastapi.testclient import TestClient  # noqa: E402
from vnbot.channels import Switches  # noqa: E402
from vnbot.jobs import JobQueue  # noqa: E402
from vnbot.service import Watchlist  # noqa: E402
from vnbot.web import create_app  # noqa: E402

import tradingagents.dataflows.symbols as _symbols  # noqa: E402

_symbols.listing_entry = lambda code: {"symbol": code, "organShortName": f"Công ty {code}"}   # the VCI listing, offline
app = create_app(SETTINGS, JobQueue(SETTINGS), Watchlist(SETTINGS), switches=Switches(SETTINGS.data_dir, {}))
w = app.state.taisan
w.market, w.site = FakeMarket(), None
web = TestClient(app, raise_server_exceptions=False)
AUTH, SAME = ("x", "s3cret-pass"), {"origin": "http://testserver"}


def post(path, data=None, **kw):
    r = web.post(path, data=data, auth=AUTH, headers=SAME, follow_redirects=False, **kw)
    from urllib.parse import unquote

    return r.status_code, unquote(r.headers.get("location", ""))


t.ok("GET /taisan needs the password", web.get("/taisan").status_code == 401)
page = web.get("/taisan", auth=AUTH)
t.ok("empty page: how to start, the tab marked", page.status_code == 200 and "Bắt đầu theo dõi tài sản" in page.text
     and 'href="/taisan" class="nav-item on" aria-current="page"' in page.text, page.status_code)
t.ok("header: weekday, date and the lunar date (24/10/2026 = 15/09 ÂL), the market's phase",
     "Thứ 7 24/10" in page.text and "15/09 ÂL" in page.text and "Thị trường nghỉ" in page.text)
nav = page.text.split('<nav aria-label="Điều hướng chính">', 1)[-1].split("</nav>", 1)[0]
t.ok("sidebar: every other page one tap away", '<nav aria-label="Điều hướng chính">' in page.text
     and all(f'href="{path}"' in nav for path in ("/", "/ai", "/thamdinh", "/caidat", "/phienban", "/zoom")), nav[:300])
t.ok("one h1 and a skip link", page.text.count("<h1") == 1 and 'class="skip" href="#main"' in page.text)
for path in ("/taisan/them", "/taisan/tien", "/taisan/khoptien", "/taisan/catlo", "/taisan/nguong", "/taisan/huy",
             "/taisan/nhap", "/taisan/nhap/xacnhan", "/taisan/nhap/huy", "/taisan/reset", "/taisan/lammoi",
             "/taisan/baocao"):
    t.ok(f"POST {path} without password -> 401", web.post(path, data={"x": "1"}).status_code == 401)
    t.ok(f"POST {path} cross-site -> 403", web.post(path, data={"x": "1"}, auth=AUTH,
                                                    headers={"origin": "https://evil.example"}).status_code == 403)
for path in ("/taisan/xuat.xlsx", "/taisan/xuat.csv"):
    t.ok(f"GET {path} needs the password", web.get(path).status_code == 401)
code, loc = post("/taisan/tien", {"kind": "DEPOSIT", "amount": "100.000.000", "day": "2026-04-13", "note": "vốn"})
t.ok("web deposit recorded, message back", code == 303 and "✅ #1 Nạp tiền 100.000.000" in loc
     and w.entries()[0].source == "web", loc)
code, loc = post("/taisan/them", {"kind": "BUY", "symbol": "fpt", "qty": "1.000", "price": "30.000", "day": "2026-04-14",
                                  "note": "<script>alert(1)</script>"})
t.ok("web buy recorded", "✅ #2 MUA FPT 1.000 cp @ 30.000" in loc, loc)
code, loc = post("/taisan/them", {"kind": "BUY", "symbol": "FPT", "qty": "10", "price": "30.5", "day": "2026-04-14"})
t.ok("web: a bad price refused, nothing recorded", "⚠️" in loc and len(w.entries()) == 2, loc)
code, loc = post("/taisan/them", {"kind": "SELL", "symbol": "FPT", "qty": "5000", "price": "31000", "day": "2026-04-20"})
t.ok("web: selling more than held refused", "⚠️" in loc and "chỉ có 1.000" in loc, loc)
code, loc = post("/taisan/them", {"kind": "SELL", "symbol": "FPT", "qty": "400", "price": "33000", "day": "2026-04-20"})
t.ok("web sell", "BÁN FPT 400 cp" in loc, loc)
code, loc = post("/taisan/them", {"kind": "BONUS", "symbol": "FPT", "qty": "60", "price": "", "day": "2026-04-21"})
t.ok("web bonus shares", "+60 cp thưởng" in loc, loc)
code, loc = post("/taisan/tien", {"kind": "DIVIDEND", "amount": "1.320.000", "symbol": "FPT", "day": "2026-05-04"})
t.ok("web cash dividend (the money received)", "Cổ tức FPT 1.320.000" in loc and w.book().dividends == 1_320_000, loc)
code, loc = post("/taisan/tien", {"kind": "DIVIDEND", "amount": "1000", "symbol": "", "day": "2026-05-04"})
t.ok("web dividend needs the symbol", "⚠️" in loc, loc)
code, loc = post("/taisan/tien", {"kind": "ADJUST", "amount": "-120.000", "day": "2026-05-04", "note": "phí"})
t.ok("web adjustment may be negative", "-120.000" in loc and w.book().adjustments == -120_000, loc)
code, loc = post("/taisan/khoptien", {"actual": "84.000.000"})
t.ok("web: cash matched to the broker", abs(w.book().cash - 84_000_000) < 1, loc)
code, loc = post("/taisan/catlo", {"symbol": "FPT", "stop": "25.000", "target": "40.000", "action": "save"})
t.ok("web levels saved", w.levels()["FPT"]["stop"] == 25_000 and w.levels()["FPT"]["target"] == 40_000, loc)
code, loc = post("/taisan/catlo", {"symbol": "FPT", "stop": "", "target": "", "action": "save"})
t.ok("web levels need the cut-loss", "⚠️" in loc, loc)
code, loc = post("/taisan/nguong", {"warn": "22,5", "danger": "35"})
t.ok("web thresholds", w.thresholds() == (22.5, 35.0), loc)
w.rebuild(through="2026-10-23")
w.refresh_analytics("2026-10-23")
page = web.get("/taisan", auth=AUTH).text
for piece in ("Chào buổi sáng", "Tổng tài sản", "Tỷ trọng vị thế",
              "Rủi ro &amp; phân bổ", 'role="tablist"', 'class="achart"', "Mã đang nắm giữ", "Công ty FPT",
              'class="spark"', "Đọc tin tức", 'href="/taisan?ma=FPT#nhatky"', "Chi tiết danh mục", "Thêm giao dịch",
              "Nhật ký giao dịch", "Nạp / rút / cổ tức / điều chỉnh", "Cắt lỗ &amp; mục tiêu", "Nhận xét &amp; chiến lược",
              "Kết quả cảnh báo bán", "Báo cáo cuối ngày", "Nhập file Excel", "Ngưỡng cảnh báo tỷ trọng", "Reset danh mục",
              'href="/taisan/xuat.xlsx"', 'href="/taisan/xuat.csv"', "Cắt lỗ 25.000", "Mục tiêu 40.000"):
    t.ok(f"page shows {piece!r}", piece in page)
t.ok("the stock card: price, change against the previous close, P&L", "GIÁ HIỆN TẠI".lower() in page.lower()
     and 'class="sk-chg' in page and '<div class="pnl"><div class="l">Lãi/lỗ</div>' in page)
t.ok("news opens outside, safely", 'target="_blank" rel="noopener noreferrer"' in page)
t.ok("the sections open from a link (#nhatky...) and close natively", '<details class="acc" id="nhatky">' in page
     and '<details class="acc" id="ghi">' in page)
t.ok("page escapes what was typed", "<script>alert(1)" not in page and "&lt;script&gt;alert(1)" in page)
pts = json.loads(__import__("html").unescape(page.split('data-points="', 1)[1].split('"', 1)[0]))
t.ok("chart data: one point per session from the first entry", pts["d"][0] == "2026-04-13" and pts["d"][-1] == "2026-10-23"
     and len(pts["d"]) == len(pts["t"]) == len(pts["v"]), {k: len(v) for k, v in pts.items()})
jr = web.get("/taisan?ma=FPT&loai=ban", auth=AUTH).text
t.ok("journal filtered by symbol and kind, the section open", '<details class="acc" id="nhatky" open>' in jr
     and jr.count('<li class="jr') == 1 and "400 cp × 33.000" in jr and "Bỏ lọc" in jr)
w.set_thresholds(10, 20)
hot = web.get("/taisan", auth=AUTH).text
t.ok("over the danger weight: the alert card and the to-do line", "Cảnh báo tỷ trọng · 1 mã" in hot
     and 'data-sym="FPT"' in hot and "Việc cần làm phiên 26/10" in hot and "⚠️ FPT chiếm" in hot)
w.set_thresholds(22.5, 35)
odd = web.get("/taisan?ma=%3Cx%3E&loai=zzz", auth=AUTH)
t.ok("odd filter values ignored", odd.status_code == 200 and "&lt;x&gt;" not in odd.text and "<x>" not in odd.text)
code, loc = post("/taisan/lammoi")
t.ok("Làm mới giá: back to the holdings", code == 303 and loc.endswith("#danhmuc") and "🔄" in loc, loc)
code, loc = post("/taisan/baocao")
t.ok("Gửi ngay without a channel says so", "⚠️" in loc and "kênh gửi" in loc, loc)
w.save_report("2026-10-23", "💼 TÀI SẢN · phiên 23/10/2026\nnội dung")
page = web.get("/taisan", auth=AUTH).text
t.ok("the last evening report on the page", "Báo cáo phiên 23/10/2026" in page and "💼 TÀI SẢN · phiên 23/10/2026" in page)
fresh = web.get("/taisan?baocao=1", auth=AUTH).text
t.ok("a fresh report on demand, its section open", '<details class="acc" id="baocao" open>' in fresh
     and "bản tính lúc này" in fresh)
undo_id = w.entries()[-1].id
code, loc = post("/taisan/huy", {"id": str(undo_id)})
t.ok("web undo of one line", loc.startswith("/taisan?msg=↩️ Đã huỷ #") and abs(w.book().cash - 84_000_000) > 1, loc)
code, loc = post("/taisan/huy", {"id": "abc"})
t.ok("web undo needs a number", "⚠️" in loc, loc)
r1 = web.get("/taisan/xuat.xlsx", auth=AUTH)
t.ok("export xlsx", r1.status_code == 200 and r1.content[:2] == b"PK" and "spreadsheetml" in r1.headers["content-type"]
     and "attachment" in r1.headers.get("content-disposition", ""), r1.headers)
r2 = web.get("/taisan/xuat.csv", auth=AUTH)
t.ok("export csv", r2.status_code == 200 and r2.text.startswith("﻿Ngày") and "attachment" in r2.headers["content-disposition"])
# The Excel upload: 1 MB with the password; 64 KB without it and on every other route.
code, loc = post("/taisan/nhap", files={"file": ("p.xlsx", old_app_file())})
t.ok("upload: a preview, nothing recorded yet", "📥 XEM TRƯỚC" in loc and tx.pending(w) is not None and len(w.entries()) > 5, loc)
page = web.get("/taisan", auth=AUTH).text
t.ok("the preview on the page, with the replace confirmation", "Xem trước nhập Excel" in page and "Thay bằng file này" in page
     and "/tsnhap" not in page.split("Xem trước nhập Excel")[1].split("</pre>")[0])
code, loc = post("/taisan/nhap/xacnhan", {"replace": "1"})
t.ok("replace needs the box ticked", "⚠️" in loc and tx.pending(w) is not None, loc)
code, loc = post("/taisan/nhap/xacnhan", {"replace": "1", "confirm": "1"})
t.ok("replace confirmed", "✅ Đã nhập 8 dòng" in loc and abs(w.book().cash - 58_784_700) < 1, loc)
for _ in range(600):
    if not w.rebuilding:
        break
    time.sleep(0.05)
t.ok("the history rebuilt in the background", not w.rebuilding and w.state().get("rebuilt_until") == "2026-10-23",
     w.state().get("rebuilt_until"))
code, loc = post("/taisan/nhap", files={"file": ("p.csv", b"a,b")})
t.ok("upload: only .xlsx", "Chỉ nhận file Excel" in loc, loc)
r3 = web.post("/taisan/nhap", files={"file": ("p.xlsx", b"PK" + b"A" * 300_000)}, headers=SAME)
t.ok("upload 300 KB without password -> 413", r3.status_code == 413, r3.status_code)
code, loc = post("/taisan/nhap", files={"file": ("p.xlsx", b"PK" + b"A" * 300_000)})
t.ok("upload 300 KB with the password is read", code == 303 and "⚠️" in loc, (code, loc))
code, loc = post("/taisan/nhap", files={"file": ("p.xlsx", b"PK" + b"A" * 1_200_000)})
t.ok("upload over 1 MB -> 413", code == 413, code)
code, loc = post("/taisan/tien", {"kind": "DEPOSIT", "amount": "1", "note": "A" * 300_000})
t.ok("other routes keep 64 KB with the password", code == 413, code)
r4 = web.post("/taisan/nhap", files={"file": ("p.xlsx", b"PK" + b"A" * 300_000)}, headers=SAME, auth=("x", "wrong"))
t.ok("a wrong password keeps 64 KB", r4.status_code == 413, r4.status_code)
code, loc = post("/taisan/reset", {"word": "ok"})
t.ok("reset needs XOA", "⚠️" in loc and w.entries(), loc)
code, loc = post("/taisan/reset", {"word": "xoa"})
t.ok("reset: empty book, the journal kept as a copy, thresholds kept",
     "🗑️ Đã xoá" in loc and not w.entries() and len(list(w.dir.glob("journal.jsonl.bak-*"))) == 2
     and w.thresholds() == (22.5, 35.0), loc)
t.ok("nav shows the Danh mục tab on other pages", 'href="/taisan"' in web.get("/", auth=AUTH).text)

# ───────────── 10. Telegram: a document, the moved commands ─────────────
from vnbot.telegram_bot import TelegramBot  # noqa: E402

tg = TelegramBot(SETTINGS, JobQueue(SETTINGS), Watchlist(SETTINGS))
tg.taisan = account("tg")
sent = []
tg.send = lambda chat_id, text: sent.append(text)
tg._download = lambda file_id: old_app_file()
msg = {"chat": {"id": 111}, "from": {"id": 111}, "text": "/tsnhap",
       "document": {"file_name": "portfolio.xlsx", "file_size": 9000, "file_id": "f1"}}
tg._handle(msg)
t.ok("Telegram: a .xlsx document gives the preview", sent and sent[-1].startswith("📥 XEM TRƯỚC NHẬP EXCEL · portfolio.xlsx"),
     sent[-1:])
tg._handle({**msg, "document": {**msg["document"], "file_name": "a.pdf"}})
t.ok("Telegram: only .xlsx", "Chỉ nhận file Excel" in sent[-1], sent[-1])
tg._handle({**msg, "document": {**msg["document"], "file_size": 3_000_000}})
t.ok("Telegram: 2 MB at most", "quá lớn" in sent[-1], sent[-1])
done_msgs = []
tg.taisan.rebuild_in_background = lambda on_done=None: done_msgs.append(on_done) or True
tg._handle({"chat": {"id": 111}, "from": {"id": 111}, "text": "/tsnhap xacnhan"})
t.ok("Telegram: confirm, then the rebuild starts", sent[-1].startswith("✅ Đã nhập 8 dòng") and len(done_msgs) == 1, sent[-1])
for old in ("/stvkhop", "/stvvon", "/stvthuong", "/stvcosan", "/stvxoa"):
    tg._handle({"chat": {"id": 111}, "from": {"id": 111}, "text": f"{old} FPT 100 30000"})
    t.ok(f"{old} points to the Tài sản commands", "/ts" in sent[-1] and sent[-1].startswith("⚠️"), sent[-1])
tg._handle({"chat": {"id": 111}, "from": {"id": 111}, "text": "/trogiup"})
t.ok("help lists the Tài sản commands, not the old ones", "/tsmua FPT 20 73891" in sent[-1] and "/stvkhop" not in sent[-1])

# ───────────── 11. the 15:30 report ─────────────
import vnbot.trader as T  # noqa: E402
from vnbot.service import DailyScheduler  # noqa: E402


class WL:
    def get(self):
        return []


bot = T.TradingBot(SETTINGS, WL())
acc = bot.taisan
acc.market, acc.site = FakeMarket(), None
t.ok("empty portfolio: no report", bot.taisan_day("2026-10-23") == "")
acc.add("DEPOSIT", "2026-10-19", amount=100_000_000)
acc.add("BUY", "2026-10-20", symbol="HPG", qty=1000, price=25_000)
acc.market.last = "2026-10-22"
t.ok("the day's bar not out: None (wait)", bot.taisan_day("2026-10-23") is None)
acc.market.last = "2026-10-23"
text = bot.taisan_day("2026-10-23")
t.ok("report of the day, history rebuilt and the close saved", text.startswith("💼 TÀI SẢN · phiên 23/10/2026")
     and acc.state()["rebuilt_until"] == "2026-10-22" and "2026-10-23" in acc.closes()["HPG"], text[:200])
t.ok("risk figures computed for the page", acc.state().get("analytics", {}).get("day") == "2026-10-23")
notes = []
sched = DailyScheduler(SETTINGS, JobQueue(SETTINGS), WL(), bot=bot, notify=notes.append)
NOW["t"] = at("2026-10-23", "15:29")
t.ok("not due before 15:30", sched._taisan_due() is None)
NOW["t"] = at("2026-10-23", "15:31")
t.ok("due at 15:30 on a session day", sched._taisan_due() == "2026-10-23")
sched._taisan_report("2026-10-23")
t.ok("sent once to the channels, stamped", len(notes) == 1 and notes[0].startswith("💼 TÀI SẢN")
     and sched._taisan_due() is None, notes)
NOW["t"] = at("2026-10-24", "15:31")
t.ok("no report on a Saturday", sched._taisan_due() is None)
NOW["t"] = at("2026-10-26", "15:31")
acc.market.last = "2026-10-23"
sched._taisan_report("2026-10-26")
t.ok("bar late: tried again later, not stamped", sched._taisan_next > time.time() and len(notes) == 1
     and sched._read(sched._taisan_stamp) == "2026-10-23")
sched._taisan_next = 0
NOW["t"] = at("2026-10-26", "18:01")
sched._taisan_report("2026-10-26")
t.ok("after 18:00 a missing bar is a closure: stamped, nothing sent", sched._read(sched._taisan_stamp) == "2026-10-26"
     and len(notes) == 1)
NOW["t"] = at("2026-10-24", "10:00")

t.done()
