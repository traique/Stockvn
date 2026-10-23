"""Excel and CSV for the Tài sản portfolio: the import (the owner's old app, or this
module's own export) and the export, in the layout of the old app.

The old app's file (3 sheets: "Vị thế hiện tại", "Lịch sử giao dịch", "Tổng quan") has
no deposits, withdrawals or cash dividends, and three faults this import corrects:

- bonus shares were written as "Bán" at price 0: read back as bonus shares;
- shares split across lots in fractions (CTS 3,947 + 1,053): the history is the source,
  so the holdings are whole shares again;
- a holding without its buy date: the date comes from the history.

Without a cash sheet the capital ("Vốn gốc") becomes one deposit on the first day, and
the difference between the file's cash and the cash the history gives becomes one
cash adjustment that day (a loss before tracking when negative). This module's own
export adds a "Tiền mặt" sheet, so a file it wrote comes back exactly.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from datetime import date as _date, datetime
from pathlib import Path

from vnbot import taisan as ts

SHEET_POSITIONS, SHEET_HISTORY, SHEET_SUMMARY, SHEET_CASH = "Vị thế hiện tại", "Lịch sử giao dịch", "Tổng quan", "Tiền mặt"
HISTORY_HEADERS = ["Ngày", "Loại", "Mã", "Số lượng", "Giá", "Giá trị", "Giá vốn BQ", "Lãi/Lỗ thực", "Ghi chú"]
POSITION_HEADERS = ["Mã", "Giá mua", "Số lượng", "Ngày mua", "Giá hiện tại", "Giá trị vốn", "Giá trị TT", "Lãi/Lỗ", "Lãi/Lỗ%"]
CASH_HEADERS = ["Ngày", "Loại", "Mã", "Số tiền", "Ghi chú"]
MAX_BYTES = 2_000_000
MAX_ROWS = 20_000
_CASH_KIND = {"nạp tiền": "DEPOSIT", "nap tien": "DEPOSIT", "nạp": "DEPOSIT", "rút tiền": "WITHDRAW", "rut tien": "WITHDRAW",
              "rút": "WITHDRAW", "cổ tức tiền": "DIVIDEND", "co tuc tien": "DIVIDEND", "cổ tức": "DIVIDEND",
              "điều chỉnh tiền": "ADJUST", "dieu chinh tien": "ADJUST", "điều chỉnh": "ADJUST"}


class ImportError_(ValueError):
    """A file that cannot be read as a portfolio (Vietnamese message)."""


def _day(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, _date):
        return value.isoformat()
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _num(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+", text):          # 56.300
        text = text.replace(".", "")
    text = text.replace(",", ".") if text.count(",") == 1 and "." not in text else text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _rows(ws) -> list[tuple]:
    out = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= MAX_ROWS:
            raise ImportError_(f"Sheet {ws.title} quá dài (hơn {MAX_ROWS} dòng).")
        out.append(tuple(row))
    return out


def _sheet(wb, name: str):
    for ws in wb.worksheets:
        if ws.title.strip().lower() == name.lower():
            return ws
    return None


def read(data: bytes) -> dict:
    """The sheets of a workbook as plain rows. ImportError_ for anything else."""
    import openpyxl

    if len(data) > MAX_BYTES:
        raise ImportError_(f"File quá lớn ({len(data) // 1024} KB, tối đa {MAX_BYTES // 1_000_000} MB).")
    if not data.startswith(b"PK"):
        raise ImportError_("Không phải file Excel .xlsx.")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any parse failure is a bad file
        raise ImportError_(f"Không đọc được file Excel: {type(exc).__name__}.") from None
    try:
        history = _sheet(wb, SHEET_HISTORY)
        if history is None:
            raise ImportError_(f"File thiếu sheet “{SHEET_HISTORY}”.")
        out = {"history": _rows(history)}
        for key, name in (("positions", SHEET_POSITIONS), ("summary", SHEET_SUMMARY), ("cash", SHEET_CASH)):
            ws = _sheet(wb, name)
            out[key] = _rows(ws) if ws is not None else None
    finally:
        wb.close()
    return out


def _header_index(header: tuple, names: list[str]) -> dict[str, int]:
    norm = [str(h or "").strip().lower() for h in header]
    out = {}
    for name in names:
        if name.lower() in norm:
            out[name] = norm.index(name.lower())
    return out


def _summary(rows) -> dict[str, float]:
    out = {}
    for row in rows or []:
        if row and row[0] and len(row) > 1:
            value = _num(row[1])
            if value is not None:
                out[str(row[0]).strip()] = value
    return out


def build(sheets: dict, *, today: str) -> dict:
    """The entries an import would record, what it checked, and whether it can be confirmed:
    {entries, lines, errors, ok, summary_file}."""
    from tradingagents.market import vn_rules

    errors, notes = [], []
    head = sheets["history"][0] if sheets["history"] else ()
    idx = _header_index(head, HISTORY_HEADERS)
    for needed in ("Ngày", "Loại", "Mã", "Số lượng", "Giá"):
        if needed not in idx:
            raise ImportError_(f"Sheet “{SHEET_HISTORY}” thiếu cột “{needed}”.")
    trades, file_pnl, bonus_fixed = [], {}, 0
    for n, row in enumerate(sheets["history"][1:], start=2):
        if not row or all(v in (None, "") for v in row):
            continue

        def cell(name, row=row):
            i = idx.get(name)
            return row[i] if i is not None and i < len(row) else None
        day, kind_raw, symbol = _day(cell("Ngày")), str(cell("Loại") or "").strip().lower(), str(cell("Mã") or "").strip().upper()
        qty, price = _num(cell("Số lượng")), _num(cell("Giá"))
        if day is None:
            errors.append(f"Dòng {n}: ngày không đọc được ({cell('Ngày')!r}).")
            continue
        if not re.fullmatch(r"[A-Z0-9]{1,10}", symbol):
            errors.append(f"Dòng {n}: mã không hợp lệ ({symbol!r}).")
            continue
        if qty is None or qty <= 0 or abs(qty - round(qty)) > 1e-6:
            errors.append(f"Dòng {n}: số lượng {qty!r} không phải số nguyên dương.")
            continue
        qty = int(round(qty))
        if kind_raw in ("thưởng cp", "thuong cp", "cổ tức cp", "co tuc cp", "thưởng", "cổ phiếu thưởng"):
            kind = "BONUS"
        elif kind_raw in ("bán", "ban", "sell") and not price:
            kind = "BONUS"                    # the old app wrote bonus shares as a sale at 0
            bonus_fixed += 1
        elif kind_raw in ("bán", "ban", "sell"):
            kind = "SELL"
        elif kind_raw in ("mua", "buy"):
            kind = "BUY"
        else:
            errors.append(f"Dòng {n}: loại giao dịch {cell('Loại')!r} không rõ.")
            continue
        if kind in ("BUY", "SELL") and (price is None or price < 100):
            errors.append(f"Dòng {n}: giá {price!r} không hợp lệ.")
            continue
        if kind in ("BUY", "SELL") and not vn_rules.is_session_day(day):
            errors.append(f"Dòng {n}: {ts.ddmmyyyy(day)} không có phiên giao dịch.")
            continue
        if day > today:
            errors.append(f"Dòng {n}: ngày {ts.ddmmyyyy(day)} ở tương lai.")
            continue
        note = str(cell("Ghi chú") or "").strip()[:200]
        trades.append({"kind": kind, "day": day, "symbol": symbol, "qty": qty,
                       "price": float(price) if kind != "BONUS" else 0.0, "note": note, "row": n})
        if kind == "SELL" and _num(cell("Lãi/Lỗ thực")) is not None:
            file_pnl[n] = _num(cell("Lãi/Lỗ thực"))
    if bonus_fixed:
        notes.append(f"Sửa {bonus_fixed} dòng “Bán … giá 0” của app cũ thành cổ phiếu thưởng.")
    if not trades:
        errors.append("Không có giao dịch nào đọc được.")

    summary = _summary(sheets.get("summary"))
    cash_rows = []
    if sheets.get("cash"):
        cidx = _header_index(sheets["cash"][0], CASH_HEADERS)
        for n, row in enumerate(sheets["cash"][1:], start=2):
            if not row or all(v in (None, "") for v in row):
                continue
            day = _day(row[cidx["Ngày"]]) if "Ngày" in cidx else None
            kind = _CASH_KIND.get(str(row[cidx["Loại"]] if "Loại" in cidx else "").strip().lower())
            amount = _num(row[cidx["Số tiền"]]) if "Số tiền" in cidx else None
            if day is None or kind is None or amount is None or amount == 0:
                errors.append(f"Sheet {SHEET_CASH} dòng {n}: không đọc được.")
                continue
            symbol = str(row[cidx["Mã"]] or "").strip().upper() if "Mã" in cidx else ""
            note = str(row[cidx["Ghi chú"]] or "").strip()[:200] if "Ghi chú" in cidx else ""
            if kind in ("DEPOSIT", "DIVIDEND", "WITHDRAW"):
                amount = abs(amount)
            cash_rows.append({"kind": kind, "day": day, "symbol": symbol, "amount": float(amount), "note": note})

    first_day = min([t["day"] for t in trades] + [c["day"] for c in cash_rows], default=today)
    entries = [dict(c) for c in cash_rows] + [dict(t) for t in trades]
    if not cash_rows:
        capital = summary.get("Vốn gốc")
        if capital and capital > 0:
            entries.append({"kind": "DEPOSIT", "day": first_day, "amount": float(capital),
                            "note": "Vốn gốc theo file Excel (ngày nạp không rõ)"})
            notes.append(f"Vốn gốc {ts.vnd(capital)} ghi là một lần nạp ngày {ts.ddmmyyyy(first_day)}.")
        else:
            errors.append("Sheet Tổng quan không có “Vốn gốc”: không biết số tiền đã nạp.")
    result = {"entries": [], "lines": [], "errors": errors, "ok": False, "summary_file": summary}
    if errors:
        return result
    probe = [_entry(i + 1, e, today) for i, e in enumerate(_ordered(entries))]
    try:
        book = ts.derive(probe)
    except ts.LedgerError as exc:
        result["errors"].append(str(exc))
        return result
    file_cash = summary.get("Tiền mặt")
    if file_cash is not None and abs(file_cash - book.cash) >= 1:
        gap = file_cash - book.cash
        label = ("Lỗ trước khi theo dõi" if gap < 0 else "Điều chỉnh tiền mặt") + " (khớp tiền mặt file Excel)"
        entries.append({"kind": "ADJUST", "day": first_day, "amount": round(gap, 2), "note": label})
        notes.append(f"{label}: {ts.signed(gap)} ghi ngày {ts.ddmmyyyy(first_day)}.")
        probe = [_entry(i + 1, e, today) for i, e in enumerate(_ordered(entries))]
        try:
            book = ts.derive(probe)
        except ts.LedgerError as exc:
            result["errors"].append(str(exc))
            return result
    result["entries"] = [dict(_entry_dict(e)) for e in probe]
    result["book"] = book
    result["notes"] = notes
    result["ok"] = True

    # Checks against the file.
    checks = []
    if "Lãi/Lỗ chốt" in summary:
        diff = book.realized - summary["Lãi/Lỗ chốt"]
        checks.append(("Lãi/lỗ đã chốt", f"{ts.signed(book.realized)} (file {ts.signed(summary['Lãi/Lỗ chốt'])})", abs(diff) < 1))
    row_of = {i + 1: e.get("row") for i, e in enumerate(_ordered(entries))}
    wrong = [row_of[r.entry.id] for r in book.rows
             if r.pnl is not None and row_of.get(r.entry.id) in file_pnl and abs(r.pnl - file_pnl[row_of[r.entry.id]]) >= 1]
    if file_pnl:
        checks.append(("Lãi/lỗ từng lệnh bán", f"{len(file_pnl) - len(wrong)}/{len(file_pnl)} lệnh khớp"
                       + (f" (lệch ở dòng {', '.join(map(str, wrong[:8]))})" if wrong else ""), not wrong))
    file_positions = {}
    if sheets.get("positions"):
        pidx = _header_index(sheets["positions"][0], POSITION_HEADERS)
        for row in sheets["positions"][1:]:
            if not row or "Mã" not in pidx or not row[pidx["Mã"]]:
                continue
            sym = str(row[pidx["Mã"]]).strip().upper()
            q = _num(row[pidx["Số lượng"]]) if "Số lượng" in pidx else None
            price = _num(row[pidx["Giá hiện tại"]]) if "Giá hiện tại" in pidx else None
            fp = file_positions.setdefault(sym, {"qty": 0.0, "price": None})
            fp["qty"] += q or 0.0
            if price:
                fp["price"] = price
        same = all(abs(round(fp["qty"]) - (book.holdings[s].qty if s in book.holdings else 0)) == 0
                   for s, fp in file_positions.items()) and set(file_positions) == set(book.holdings)
        checks.append(("Số cổ phiếu đang giữ", ", ".join(f"{s} {ts.vnd(h.qty)}" for s, h in sorted(book.holdings.items()))
                       or "không giữ mã nào", same))
    prices = {s: fp["price"] for s, fp in file_positions.items() if fp.get("price")}
    if prices and "Lãi/Lỗ tổng" in summary:
        market = sum(h.qty * (prices.get(s) or h.avg) for s, h in book.holdings.items())
        total = book.cash + market
        diff = total - book.capital - summary["Lãi/Lỗ tổng"]
        checks.append(("Lãi/lỗ tổng theo giá trong file", f"{ts.signed(total - book.capital)} (file {ts.signed(summary['Lãi/Lỗ tổng'])})",
                        abs(diff) < 1))
    if file_cash is not None:
        checks.append(("Tiền mặt", f"{ts.vnd(book.cash)} (file {ts.vnd(file_cash)})", abs(book.cash - file_cash) < 1))
    result["checks"] = checks
    result["prices"] = prices
    return result


def _ordered(entries: list[dict]) -> list[dict]:
    return sorted(entries, key=lambda e: (e["day"], ts._RANK.get(e["kind"], 9), e.get("row", 0)))


def _entry(i: int, e: dict, today: str) -> ts.Entry:
    return ts.Entry(i, e["kind"], e["day"], e.get("symbol", ""), int(e.get("qty") or 0), float(e.get("price") or 0),
                    float(e.get("amount") or 0), e.get("note", ""), at=f"{today}T00:00:00", source="import")


def _entry_dict(e: ts.Entry) -> dict:
    from dataclasses import asdict

    return asdict(e)


def preview(account: ts.TaiSan, data: bytes, *, filename: str = "") -> str:
    """Read ``data``, keep what the import would record (``import_pending.json``) and say it."""
    sheets = read(data)
    today = ts.today_vn()
    got = build(sheets, today=today)
    if not got["ok"]:
        account.import_path.unlink(missing_ok=True)
        return ("❌ Chưa nhập được file" + (f" {filename}" if filename else "") + ":\n"
                + "\n".join(f"• {e}" for e in got["errors"][:15]))
    book = got["book"]
    entries = [ts.Entry.from_dict(e) for e in got["entries"]]
    kinds = {}
    for e in entries:
        kinds[e.kind] = kinds.get(e.kind, 0) + 1
    current = ts.effective(account.entries())
    digest = hashlib.sha256(data).hexdigest()[:12]
    lines = [f"📥 XEM TRƯỚC NHẬP EXCEL{f' · {filename}' if filename else ''}",
             f"{len(entries)} dòng: " + ", ".join(f"{ts.KIND_VI[k]} {n}" for k, n in sorted(kinds.items())),
             f"Từ {ts.ddmmyyyy(book.first_day)} tới {ts.ddmmyyyy(book.last_day)}", ""]
    lines += [f"• {n}" for n in got.get("notes") or []]
    lines += ["", "Danh mục sau khi nhập:"]
    lines += [f"• {s} {ts.vnd(h.qty)} cp · giá vốn BQ {ts.vnd(h.avg)}" for s, h in sorted(book.holdings.items())] or ["• không giữ mã nào"]
    lines += [f"Tiền mặt {ts.vnd(book.cash)} · vốn gốc {ts.vnd(book.capital)}",
              f"Đã chốt {ts.signed(book.realized)} ({book.sells} lệnh bán, thắng {book.wins}, thua {book.losses})"]
    checks = got.get("checks") or []
    if checks:
        lines += ["", "Đối chiếu với file:"] + [f"{'✅' if ok else '⚠️'} {label}: {text}" for label, text, ok in checks]
    lines += [f"⚠️ {w}" for w in book.warnings]
    lines.append("")
    if current:
        lines.append(f"Danh mục Tài sản đang có {len(current)} dòng. Thay toàn bộ bằng file này: /tsnhap thaythe "
                     "(sổ cũ được lưu bản sao). Bỏ: /tsnhap huy")
    else:
        lines.append("Ghi vào danh mục: /tsnhap xacnhan · bỏ: /tsnhap huy")
    text = "\n".join(lines)
    ts._write_json(account.import_path, {"created": ts.now_vn().isoformat(timespec="seconds"), "digest": digest,
                                         "filename": filename, "entries": got["entries"],
                                         "prices": got.get("prices") or {}, "text": text})
    return text


def pending(account: ts.TaiSan) -> dict | None:
    """The file waiting for confirmation (its preview ``text``, ``filename``, ``created``), or None."""
    data = ts._read_json(account.import_path, {})
    return data if isinstance(data, dict) and data.get("entries") else None


def confirm(account: ts.TaiSan, *, replace: bool) -> str:
    pending = ts._read_json(account.import_path, {})
    if not pending.get("entries"):
        raise ts.LedgerError("Chưa có file nào chờ nhập: gửi file .xlsx trước (Telegram hoặc trang /taisan).")
    entries = [ts.Entry.from_dict(e) for e in pending["entries"]]
    with account._lock:
        current = ts.effective(account.entries())
        if current and not replace:
            raise ts.LedgerError(f"Danh mục đang có {len(current)} dòng: dùng /tsnhap thaythe để thay toàn bộ "
                                 "(sổ cũ được lưu bản sao).")
        ts.derive(entries)
        if account.journal_path.exists():
            account.backup_journal()
        account._append(entries)
        prices = pending.get("prices") or {}

        def reset(s):
            s["levels"], s["marks"] = {}, {}
            s["rebuilt_until"] = ""
            s["imported"] = {"at": ts.now_vn().isoformat(timespec="seconds"), "digest": pending.get("digest"),
                             "filename": pending.get("filename"), "file_prices": prices}
        account._update_state(reset)
        account.import_path.unlink(missing_ok=True)
    account._forget_live()
    book = account.book()
    return (f"✅ Đã nhập {len(entries)} dòng vào danh mục Tài sản: {len(book.holdings)} mã, tiền mặt {ts.vnd(book.cash)}, "
            f"vốn gốc {ts.vnd(book.capital)}. Bot đang dựng lại biểu đồ tài sản từ {ts.ddmmyyyy(book.first_day)} "
            "(vài phút). Xem: /ts hoặc trang /taisan.")


def cancel(account: ts.TaiSan) -> str:
    existed = account.import_path.exists()
    account.import_path.unlink(missing_ok=True)
    return "Đã bỏ file chờ nhập." if existed else "Không có file nào chờ nhập."


# ─── export ──────────────────────────────────────────────────────────────────

def _position_rows(summary: dict) -> list[list]:
    rows = [POSITION_HEADERS]
    for p in summary["positions"]:
        rows.append([p["symbol"], round(p["avg"], 2), p["qty"], p["first_day"], p["price"] or "",
                     round(p["cost"]), round(p["value"]) if p["value"] is not None else "",
                     round(p["pnl"]) if p["pnl"] is not None else "",
                     round(p["pnl_pct"], 2) if p["pnl_pct"] is not None else ""])
    return rows


def _history_rows(book: ts.Book) -> list[list]:
    rows = [HISTORY_HEADERS]
    stock = [r for r in book.rows if r.entry.kind in ts.STOCK_KINDS]
    for r in sorted(stock, key=lambda r: (r.entry.day, r.entry.id), reverse=True):
        e = r.entry
        rows.append([e.day, ts.KIND_VI[e.kind], e.symbol, e.qty, e.price if e.kind != "BONUS" else 0,
                     round(e.qty * e.price), round(r.avg_cost, 2) if r.avg_cost is not None else "",
                     round(r.pnl, 2) if r.pnl is not None else "", e.note])
    return rows


def _cash_rows(book: ts.Book) -> list[list]:
    rows = [CASH_HEADERS]
    for r in book.rows:
        e = r.entry
        if e.kind in ts.CASH_KINDS:
            amount = -e.amount if e.kind == "WITHDRAW" else e.amount
            rows.append([e.day, ts.KIND_VI[e.kind], e.symbol, round(amount, 2), e.note])
    return rows


def _summary_rows(summary: dict, book: ts.Book) -> list[list]:
    win = f"{book.wins / book.sells * 100:.1f}%" if book.sells else "—"
    return [["TỔNG QUAN DANH MỤC", ""],
            ["Tổng tài sản (NAV)", round(summary["total"])],
            ["Giá trị cổ phiếu", round(summary["market"])],
            ["Tiền mặt", round(summary["cash"])],
            ["Quyền chờ về", round(summary["pending"])],
            ["Vốn gốc", round(summary["capital"])],
            ["Lãi/Lỗ tổng", round(summary["pnl"])],
            ["Lãi/Lỗ chưa thực hiện", round(summary["unrealized"])],
            ["", ""],
            ["GIAO DỊCH ĐÃ CHỐT", ""],
            ["Số lần bán", book.sells],
            ["Lãi/Lỗ chốt", round(book.realized, 2)],
            ["Win rate", win],
            ["Thắng / Thua", f"{book.wins} / {book.losses}"],
            ["Cổ tức đã nhận", round(book.dividends)],
            ["", ""],
            ["Số mã đang nắm", len(summary["positions"])],
            ["Ngày xuất", ts.now_vn().strftime("%d/%m/%Y")]]


def export_xlsx(account: ts.TaiSan, path: Path) -> Path:
    """The four sheets, at the current prices."""
    import openpyxl
    from openpyxl.styles import Font

    book = account.book()
    summary = account.summary(book=book)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in ((SHEET_POSITIONS, _position_rows(summary)), (SHEET_HISTORY, _history_rows(book)),
                       (SHEET_SUMMARY, _summary_rows(summary, book)), (SHEET_CASH, _cash_rows(book))):
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 48)
            for c in col[1:]:
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    c.number_format = "#,##0" if float(c.value).is_integer() else "#,##0.00"
                elif isinstance(c.value, str) and c.value.startswith("="):
                    c.data_type = "s"            # a note typed as "=..." stays text, never a formula
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_csv(account: ts.TaiSan) -> str:
    """Every entry, oldest first (a BOM so Excel reads the Vietnamese)."""
    book = account.book()
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["Ngày", "Loại", "Mã", "Số lượng", "Giá", "Số tiền", "Giá vốn BQ", "Lãi/Lỗ thực", "Ghi chú"])
    for r in book.rows:
        e = r.entry
        amount = (-e.amount if e.kind == "WITHDRAW" else e.amount) if e.kind in ts.CASH_KINDS else round(e.qty * e.price)
        writer.writerow([e.day, ts.KIND_VI[e.kind], e.symbol, e.qty or "", e.price or "", round(amount, 2),
                         round(r.avg_cost, 2) if r.avg_cost is not None else "",
                         round(r.pnl, 2) if r.pnl is not None else "", _cell_text(e.note)])
    return "\ufeff" + out.getvalue()


def _cell_text(text: str) -> str:
    """A note a spreadsheet would read as a formula (=, +, -, @) is kept as text."""
    return "'" + text if text[:1] in ("=", "+", "-", "@") else text


def json_dump(value) -> str:
    return json.dumps(value, ensure_ascii=False)
